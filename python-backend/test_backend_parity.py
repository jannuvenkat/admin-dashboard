import io
import openpyxl
from starlette.testclient import TestClient
from app.main import app

def run_all_tests():
    print("==================================================")
    print("  Starting Python FastAPI Backend Parity Tests")
    print("==================================================")

    # Use TestClient with lifespan context manager
    with TestClient(app) as client:
        # 1. Health check
        print("\n[Test 1] Health check: GET /api/health")
        res = client.get("/api/health")
        assert res.status_code == 200, f"Expected 200, got {res.status_code}"
        body = res.json()
        assert body["success"] is True
        print("  -> PASSED: Health check successful")

        # 2. Login
        print("\n[Test 2] Admin Login: POST /api/auth/login")
        res = client.post("/api/auth/login", json={
            "email": "admin@example.com",
            "password": "Admin@12345"
        })
        assert res.status_code == 200, f"Login failed: {res.text}"
        body = res.json()
        assert body["success"] is True
        token = body["data"]["token"]
        user = body["data"]["user"]
        assert token and len(token) > 20
        assert user["role"] == "admin"
        assert user["email"] == "admin@example.com"
        auth_headers = {"Authorization": f"Bearer {token}"}
        print(f"  -> PASSED: Token acquired for admin ({user['email']})")

        # 3. Invalid Login
        print("\n[Test 3] Invalid Login: POST /api/auth/login")
        res = client.post("/api/auth/login", json={
            "email": "admin@example.com",
            "password": "WrongPassword999"
        })
        assert res.status_code == 401
        assert res.json()["success"] is False
        print("  -> PASSED: Rejected invalid credentials with 401")

        # 4. Auth Me
        print("\n[Test 4] Profile: GET /api/auth/me")
        res = client.get("/api/auth/me", headers=auth_headers)
        assert res.status_code == 200
        assert res.json()["data"]["email"] == "admin@example.com"

        res_no_auth = client.get("/api/auth/me")
        assert res_no_auth.status_code == 401
        print("  -> PASSED: Auth verification & unauthorized rejection work")

        # 5. Logout
        print("\n[Test 5] Logout: POST /api/auth/logout")
        res = client.post("/api/auth/logout")
        assert res.status_code == 200
        assert res.json()["success"] is True
        print("  -> PASSED: Logout successful")

        # 6. Dashboard Stats
        print("\n[Test 6] Dashboard Stats: GET /api/dashboard/stats")
        res = client.get("/api/dashboard/stats", headers=auth_headers)
        assert res.status_code == 200
        data = res.json()["data"]
        assert "totalUsers" in data and data["totalUsers"] > 0
        assert "totalProducts" in data and data["totalProducts"] > 0
        assert "totalOrders" in data and data["totalOrders"] > 0
        assert "totalRevenue" in data and data["totalRevenue"] > 0
        assert "changes" in data
        assert "recentOrders" in data and len(data["recentOrders"]) > 0
        print(f"  -> PASSED: Dashboard returned {data['totalUsers']} users, {data['totalOrders']} orders, ${data['totalRevenue']} revenue")

        # 7. Analytics: Revenue
        print("\n[Test 7] Analytics Revenue: GET /api/analytics/revenue")
        res = client.get("/api/analytics/revenue", headers=auth_headers)
        assert res.status_code == 200
        data = res.json()["data"]
        assert "totalRevenue" in data
        assert "daily" in data and len(data["daily"]) > 0
        assert "monthly" in data and len(data["monthly"]) > 0
        print("  -> PASSED: Revenue analytics breakdown returned")

        # 8. Analytics: Users
        print("\n[Test 8] Analytics Users: GET /api/analytics/users")
        res = client.get("/api/analytics/users", headers=auth_headers)
        assert res.status_code == 200
        data = res.json()["data"]
        assert "totalUsers" in data
        assert "daily" in data
        print("  -> PASSED: User analytics returned")

        # 9. Analytics: Orders
        print("\n[Test 9] Analytics Orders: GET /api/analytics/orders")
        res = client.get("/api/analytics/orders", headers=auth_headers)
        assert res.status_code == 200
        data = res.json()["data"]
        assert "totalOrders" in data
        assert "statusBreakdown" in data and len(data["statusBreakdown"]) == 4
        print("  -> PASSED: Orders analytics and status breakdown returned")

        # 10. Analytics: Products
        print("\n[Test 10] Analytics Products: GET /api/analytics/products")
        res = client.get("/api/analytics/products", headers=auth_headers)
        assert res.status_code == 200
        data = res.json()["data"]
        assert "topSelling" in data and len(data["topSelling"]) > 0
        assert "categoryRevenue" in data and len(data["categoryRevenue"]) > 0
        print("  -> PASSED: Products analytics returned")

        # 11. Users Management
        print("\n[Test 11] Users List: GET /api/users")
        res = client.get("/api/users?page=1&limit=5", headers=auth_headers)
        assert res.status_code == 200
        body = res.json()
        assert len(body["data"]) == 5
        assert body["pagination"]["total"] >= 26
        print(f"  -> PASSED: Users list returned with pagination total {body['pagination']['total']}")

        # 12. Single User
        print("\n[Test 12] Single User: GET /api/users/1")
        res = client.get("/api/users/1", headers=auth_headers)
        assert res.status_code == 200
        assert res.json()["data"]["id"] == 1
        print("  -> PASSED: Single user fetched")

        # 13. Products Management
        print("\n[Test 13] Products List: GET /api/products")
        res = client.get("/api/products?page=1&limit=5", headers=auth_headers)
        assert res.status_code == 200
        body = res.json()
        assert len(body["data"]) == 5
        assert body["pagination"]["total"] >= 35
        print(f"  -> PASSED: Products list returned with pagination total {body['pagination']['total']}")

        # 14. Product Categories
        print("\n[Test 14] Product Categories: GET /api/products/categories")
        res = client.get("/api/products/categories", headers=auth_headers)
        assert res.status_code == 200
        cats = res.json()["data"]
        assert len(cats) >= 6
        print(f"  -> PASSED: {len(cats)} categories returned")

        # 15. Single Product
        print("\n[Test 15] Single Product: GET /api/products/1")
        res = client.get("/api/products/1", headers=auth_headers)
        assert res.status_code == 200
        assert res.json()["data"]["id"] == 1
        print("  -> PASSED: Single product fetched")

        # 16. Orders Management
        print("\n[Test 16] Orders List: GET /api/orders")
        res = client.get("/api/orders?page=1&limit=5", headers=auth_headers)
        assert res.status_code == 200
        body = res.json()
        assert len(body["data"]) == 5
        assert body["pagination"]["total"] >= 65
        print(f"  -> PASSED: Orders list returned with pagination total {body['pagination']['total']}")

        # 17. Single Order with Items
        print("\n[Test 17] Single Order: GET /api/orders/1")
        res = client.get("/api/orders/1", headers=auth_headers)
        assert res.status_code == 200
        order = res.json()["data"]
        assert order["id"] == 1
        assert "items" in order and len(order["items"]) > 0
        print(f"  -> PASSED: Order #1 fetched with {len(order['items'])} items")

        # 18. Order Status Update
        print("\n[Test 18] Order Status Update: PUT /api/orders/1/status")
        orig_status = order["status"]
        test_status = "processing" if orig_status != "processing" else "pending"
        res = client.put("/api/orders/1/status", json={"status": test_status}, headers=auth_headers)
        assert res.status_code == 200
        assert res.json()["data"]["status"] == test_status
        # Restore original status
        res = client.put("/api/orders/1/status", json={"status": orig_status}, headers=auth_headers)
        assert res.status_code == 200
        assert res.json()["data"]["status"] == orig_status
        print(f"  -> PASSED: Order status updated to {test_status} and safely restored to {orig_status}")

        # 19. Excel Export (No filters)
        print("\n[Test 19] Excel Export: GET /api/orders/export")
        res = client.get("/api/orders/export", headers=auth_headers)
        assert res.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
        disposition = res.headers.get("content-disposition", "")
        assert "orders_" in disposition and ".xlsx" in disposition
        # Verify valid XLSX structure with openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        assert "Orders" in wb.sheetnames
        ws = wb["Orders"]
        assert ws.cell(row=1, column=1).value == "Orders Report"
        assert ws.cell(row=5, column=1).value == "Order ID"
        assert ws.cell(row=5, column=8).value == "Order Date"
        assert ws.max_row >= 65
        print(f"  -> PASSED: Full Excel file generated ({len(res.content)} bytes, {ws.max_row} rows, worksheet 'Orders')")

        # 20. Excel Export with Filters
        print("\n[Test 20] Excel Export with Filters: GET /api/orders/export?status=pending")
        res = client.get("/api/orders/export?status=pending", headers=auth_headers)
        assert res.status_code == 200
        wb_filt = openpyxl.load_workbook(io.BytesIO(res.content))
        ws_filt = wb_filt["Orders"]
        filter_text = ws_filt.cell(row=3, column=1).value
        assert "Status: Pending" in filter_text
        print("  -> PASSED: Filtered Excel export contains 'Status: Pending' filter tag")

        # 21. Excel Export 404 for empty results
        print("\n[Test 21] Excel Export Empty Results: GET /api/orders/export?search=NonExistentUser9999")
        res = client.get("/api/orders/export?search=NonExistentUser9999", headers=auth_headers)
        assert res.status_code == 404
        body = res.json()
        assert body["success"] is False
        assert "No orders found" in body["message"]
        print("  -> PASSED: Empty export returned 404 with exact error message")

        # 22. User CRUD Parity
        print("\n[Test 22] User CRUD Parity: POST /api/users -> PUT -> DELETE")
        create_res = client.post("/api/users", json={
            "name": "Parity Test User",
            "email": "parity_test_user@example.com",
            "password": "Password@123",
            "role": "user"
        }, headers=auth_headers)
        assert create_res.status_code == 201
        created_user = create_res.json()["data"]
        uid = created_user["id"]
        assert created_user["name"] == "Parity Test User"

        update_res = client.put(f"/api/users/{uid}", json={
            "name": "Parity Test User Updated"
        }, headers=auth_headers)
        assert update_res.status_code == 200
        assert update_res.json()["data"]["name"] == "Parity Test User Updated"

        del_res = client.delete(f"/api/users/{uid}", headers=auth_headers)
        assert del_res.status_code == 200
        assert del_res.json()["success"] is True

        # Verify deleted
        get_res = client.get(f"/api/users/{uid}", headers=auth_headers)
        assert get_res.status_code == 404
        print(f"  -> PASSED: User created (id={uid}), updated, and deleted cleanly")

        # 23. Product CRUD Parity
        print("\n[Test 23] Product CRUD Parity: POST /api/products -> PUT -> DELETE")
        prod_res = client.post("/api/products", json={
            "name": "Parity Test Product",
            "category_id": 1,
            "price": 89.99,
            "stock": 25,
            "status": "active"
        }, headers=auth_headers)
        assert prod_res.status_code == 201
        created_prod = prod_res.json()["data"]
        pid = created_prod["id"]

        upd_prod = client.put(f"/api/products/{pid}", json={
            "price": 99.99
        }, headers=auth_headers)
        assert upd_prod.status_code == 200
        assert upd_prod.json()["data"]["price"] == 99.99

        del_prod = client.delete(f"/api/products/{pid}", headers=auth_headers)
        assert del_prod.status_code == 200
        assert del_prod.json()["success"] is True

        get_prod = client.get(f"/api/products/{pid}", headers=auth_headers)
        assert get_prod.status_code == 404
        print(f"  -> PASSED: Product created (id={pid}), updated, and deleted cleanly")

    print("\n==================================================")
    print("  ALL 23 TESTS PASSED SUCCESSFULLY! 100% PARITY!  ")
    print("==================================================")

if __name__ == "__main__":
    run_all_tests()
