import io
from datetime import datetime, date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def format_date(date_val) -> str:
    if not date_val:
        return "N/A"
    if isinstance(date_val, str):
        try:
            if "T" in date_val:
                dt = datetime.fromisoformat(date_val)
            elif " " in date_val:
                dt = datetime.strptime(date_val.split(".")[0], "%Y-%m-%d %H:%M:%S")
            else:
                dt = datetime.strptime(date_val, "%Y-%m-%d")
            return dt.strftime("%d-%b-%Y %H:%M")
        except Exception:
            return date_val
    if isinstance(date_val, datetime):
        return date_val.strftime("%d-%b-%Y %H:%M")
    if isinstance(date_val, date):
        return date_val.strftime("%d-%b-%Y")
    return str(date_val)

class ExcelService:
    @staticmethod
    def generate_orders_workbook(orders: list[dict], filters: dict) -> io.BytesIO:
        search = filters.get("search", "")
        status = filters.get("status", "")
        start_date = filters.get("startDate", "")
        end_date = filters.get("endDate", "")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Orders"
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.page_setup.fitToPage = True

        # Styles
        title_font = Font(name="Calibri", size=16, bold=True, color="1E293B")
        gen_font = Font(name="Calibri", size=10, italic=True, color="64748B")
        filter_font = Font(name="Calibri", size=10, color="475569")

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_border = Border(
            top=Side(style="thin", color="0F172A"),
            bottom=Side(style="medium", color="0F172A"),
            left=Side(style="thin", color="334155"),
            right=Side(style="thin", color="334155")
        )

        data_font = Font(name="Calibri", size=10)
        data_border = Border(
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0")
        )
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        total_font = Font(name="Calibri", size=11, bold=True, color="0F172A")

        # 1. Report Title
        ws.append(["Orders Report"])
        ws.row_dimensions[1].height = 28
        ws.cell(row=1, column=1).font = title_font

        # 2. Generation Timestamp
        gen_str = format_date(datetime.now())
        ws.append([f"Generated: {gen_str}"])
        ws.cell(row=2, column=1).font = gen_font

        # 3. Active Filters
        filter_parts = []
        if search:
            filter_parts.append(f"Search: {search}")
        if status:
            filter_parts.append(f"Status: {status.capitalize()}")
        if start_date:
            filter_parts.append(f"From: {start_date}")
        if end_date:
            filter_parts.append(f"To: {end_date}")

        filter_text = f"Filters: {' | '.join(filter_parts)}" if filter_parts else "Filters: All Orders (No filters applied)"
        ws.append([filter_text])
        ws.cell(row=3, column=1).font = filter_font

        # 4. Blank spacing row
        ws.append([])

        # 5. Table Header Row (Row 5)
        headers = [
            "Order ID",
            "Customer Name",
            "Customer Email",
            "Amount",
            "Order Status",
            "Payment Method",
            "Payment Status",
            "Order Date"
        ]
        ws.append(headers)
        header_row_index = 5
        ws.row_dimensions[header_row_index].height = 26

        for col_idx in range(1, 9):
            cell = ws.cell(row=header_row_index, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = header_border

        # 6. Data Rows
        total_revenue = 0.0
        for index, order in enumerate(orders):
            curr_row = header_row_index + 1 + index
            ws.row_dimensions[curr_row].height = 20

            amount_val = float(order.get("total_amount") or 0)
            total_revenue += amount_val

            row_data = [
                f"#{order.get('id')}",
                order.get("customer_name") or "N/A",
                order.get("customer_email") or "N/A",
                amount_val,
                (order.get("status") or "pending").upper(),
                order.get("payment_method") or "Online",
                (order.get("payment_status") or "pending").upper(),
                format_date(order.get("created_at"))
            ]
            ws.append(row_data)

            # Cell formatting
            alignments = [
                Alignment(horizontal="center", vertical="center"),
                Alignment(horizontal="left", vertical="center"),
                Alignment(horizontal="left", vertical="center"),
                Alignment(horizontal="right", vertical="center"),
                Alignment(horizontal="center", vertical="center"),
                Alignment(horizontal="center", vertical="center"),
                Alignment(horizontal="center", vertical="center"),
                Alignment(horizontal="center", vertical="center"),
            ]

            is_even = (index % 2 == 0)
            for col_idx in range(1, 9):
                cell = ws.cell(row=curr_row, column=col_idx)
                cell.font = data_font
                cell.alignment = alignments[col_idx - 1]
                cell.border = data_border
                if not is_even:
                    cell.fill = zebra_fill

            # Format amount as currency
            ws.cell(row=curr_row, column=4).number_format = '"$"#,##0.00'

        # 7. Totals Section
        ws.append([])  # Blank row
        total_orders_row = header_row_index + len(orders) + 2
        ws.append(["Total Orders:", len(orders)])
        ws.cell(row=total_orders_row, column=1).font = total_font
        ws.cell(row=total_orders_row, column=2).font = total_font
        ws.cell(row=total_orders_row, column=1).alignment = Alignment(horizontal="left")
        ws.cell(row=total_orders_row, column=2).alignment = Alignment(horizontal="left")

        total_rev_row = total_orders_row + 1
        ws.append(["Total Revenue:", total_revenue])
        ws.cell(row=total_rev_row, column=1).font = total_font
        ws.cell(row=total_rev_row, column=2).font = total_font
        ws.cell(row=total_rev_row, column=1).alignment = Alignment(horizontal="left")
        ws.cell(row=total_rev_row, column=2).alignment = Alignment(horizontal="left")
        ws.cell(row=total_rev_row, column=2).number_format = '"$"#,##0.00'

        # 8. Auto-filter
        ws.auto_filter.ref = f"A{header_row_index}:H{header_row_index}"

        # 9. Freeze panes below header
        ws.freeze_panes = f"A{header_row_index + 1}"

        # 10. Auto-fit column widths
        for col_idx in range(1, 9):
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for row_idx in range(header_row_index, header_row_index + len(orders) + 1):
                val = str(ws.cell(row=row_idx, column=col_idx).value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 35)

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

excel_service = ExcelService()
